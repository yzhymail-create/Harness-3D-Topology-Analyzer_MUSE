#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""线束拓扑分析报告 Excel 生成. make_report(result, xlsx_path)

命名/编码规则(2026-09-27 方案):
- 线段: 全局 SEG01..(主干被搭接/固定站位切分后的每一段)
- 节点: CON01..(连接器, 只连一个线段端头, 3D实物名)/ CLP01..(固定卡扣, 路径中间, 3D实物名)/
         BN01..(自动分支点, 度数>=3)/ N01..(普通连接点/悬空端)
- 编码对照表: 编码 <-> 3D实例全名, 编码为稳定主键, 3D名替换时只改对照表
- 线段表的起点/终点为节点编码, 同名即相连, 以此表达拓扑连接关系
- 提取端只记录 3D 真值(坐标/弧长/直径), 不做 2D 展开
"""
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# CATIA 实例全名映射(用户确认, 2026-09-27)
# STEP标签 -> CATIA产品树实例全名
INSTANCE_MAP = {
    # 10条管状分支 -> HAR\Multi-branchable
    "S12": "HAR\\Multi-branchable1.1",
    "S13": "HAR\\Multi-branchable2.1",
    "S14": "HAR\\Multi-branchable3.1",
    "S15": "HAR\\Multi-branchable4.1",
    "S16": "HAR\\Multi-branchable5.1",
    "S17": "HAR\\Multi-branchable6.1",
    "S18": "HAR\\Multi-branchable7.1",
    "S19": "HAR\\Multi-branchable8.1",
    "S20": "HAR\\Multi-branchable9.1",
    "S21": "HAR\\Multi-branchable10.1",
    # 7个连接器 -> CONNECTOR\<名>\test
    "S4": "CONNECTOR\\测试1\\test",
    "S5": "CONNECTOR\\AXA\\test",
    "S6": "CONNECTOR\\ABU\\test",
    "S7": "CONNECTOR\\BMS\\test",
    "S8": "CONNECTOR\\BMS\\test",
    "S9": "CONNECTOR\\测试2\\test",
    "S10": "CONNECTOR\\测试3\\test",
    "S11": "CONNECTOR\\测试4\\test",
    # 4个固定卡扣 -> SUPPORT\CP0?\Tyton-157-00181 (待用户确认CP编号与S0-S3对应关系)
}

def instance_name(step_tag):
    """STEP标签(如 X1X-DRD_AllCATPart#S12) -> CATIA实例全名."""
    if not step_tag:
        return ""
    stag = step_tag.split("#")[-1]
    # 处理合并标签 S7+S8
    if "+" in stag:
        parts = []
        for t in stag.split("+"):
            t = t if t.startswith("S") else "S" + t
            parts.append(INSTANCE_MAP.get(t, t))
        # 去重
        seen = []
        for pn in parts:
            if pn not in seen:
                seen.append(pn)
        return "; ".join(seen)
    return INSTANCE_MAP.get(stag, step_tag)

HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="4472C4")
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

TYPE_CN = {"connector": "连接器", "clamp": "固定卡扣", "tie": "扎带", "fork": "分支点",
           "hanging": "悬空端", "joint": "连接点"}


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = HDR; c.fill = FILL; c.alignment = CENTER; c.border = BORDER
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER; c.alignment = CENTER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def _chain_runs(segments, nodes):
    """走线串连: 在度数!=2 的节点处截断(与新版构建器一致)."""
    nmap = {nd["code"]: nd for nd in nodes}
    adj = {nd["code"]: [] for nd in nodes}
    smap = {s["seg"]: s for s in segments}
    for s in segments:
        adj[s["na"]].append(s["seg"])
        adj[s["nb"]].append(s["seg"])

    def other(sg, code):
        return sg["nb"] if sg["na"] == code else sg["na"]

    unseen = set(smap)
    seq = []

    def walk(start_code, seg):
        path = []
        cur_code, cur_seg = start_code, seg
        while True:
            unseen.discard(cur_seg)
            path.append(cur_seg)
            nxt = other(smap[cur_seg], cur_code)
            if nmap[nxt]["degree"] != 2:
                return path, start_code, nxt
            cands = [x for x in adj[nxt] if x in unseen]
            if not cands:
                return path, start_code, nxt
            cur_seg, cur_code = cands[0], nxt

    for nd in sorted(nodes, key=lambda r: r["id"]):
        if nd["degree"] == 2:
            continue
        for sg in list(adj[nd["code"]]):
            if sg in unseen:
                path, na, nb = walk(nd["code"], sg)
                seq.append((path, na, nb))
    while unseen:
        sg = next(iter(unseen))
        path, na, nb = walk(smap[sg]["na"], sg)
        seq.append((path, na, nb))
    out = []
    for i, (path, na, nb) in enumerate(seq, 1):
        L = sum(smap[s]["length"] for s in path)
        out.append({"run": f"走线{i}", "segs": path, "length": round(L, 1),
                    "na": na, "nb": nb})
    return out


def _adapt_old(d):
    """旧版 topology.json(无 segments 键) -> 新版结构, 保证旧数据仍可出表."""
    if d.get("segments"):
        return d
    branches, nodes = d["branches"], d["nodes"]
    TOL = d["node_tol"]
    code_of, c_n, c_bn = {}, 0, 0
    for nd in nodes:
        if nd["degree"] >= 3:
            c_bn += 1; code = f"BN{c_bn:02d}"
        else:
            c_n += 1; code = f"N{c_n:02d}"
        code_of[nd["id"]] = code

    def node_of(pt):
        best, bd = None, 1e9
        for nd in nodes:
            dd = float(np.linalg.norm(np.array(nd["xyz"]) - np.array(pt)))
            if dd < bd:
                bd, best = dd, nd["id"]
        return best if bd <= TOL + 1e-6 else None

    segments, runs, entities = [], [], []
    for i, b in enumerate(branches, 1):
        seg = f"SEG{i:02d}"
        na, nb = node_of(b["p0"]), node_of(b["p1"])
        ca, cb = code_of.get(na, "?"), code_of.get(nb, "?")
        segments.append({"seg": seg, "branch": b.get("key", b["proto"]), "tube": b["occ"],
                         "na": ca, "nb": cb, "length": b["length"], "dia": b["dia"],
                         "p0": b["p0"], "p1": b["p1"], "run": f"走线{i}"})
        runs.append({"run": f"走线{i}", "segs": [seg], "length": b["length"],
                     "na": ca, "nb": cb})
        entities.append({"code": seg, "category": "线段", "name_3d": b["occ"],
                         "note": f"对应分支{b.get('key', b['proto'])}"})
        b["segs"] = [seg]; b["node0"] = ca; b["node1"] = cb
    nodes2 = []
    for nd in nodes:
        code = code_of[nd["id"]]
        ntype = "fork" if nd["degree"] >= 3 else ("hanging" if nd["degree"] == 1 else "joint")
        segs = [s["seg"] for s in segments if s["na"] == code or s["nb"] == code]
        nearest = None
        if nd.get("nearest_solid"):
            nearest = {"tag": nd["nearest_solid"], "proto": nd.get("nearest_proto"),
                       "dist": nd.get("dist")}
        nodes2.append({"id": nd["id"], "code": code, "name_3d": "", "type": ntype,
                       "xyz": nd["xyz"], "degree": nd["degree"], "segments": segs,
                       "entity": None, "nearest": nearest})
        entities.append({"code": code, "category": "分支点" if ntype == "fork" else "节点",
                         "name_3d": "", "note": "旧版数据适配, 无接触分析"})
    runs = _chain_runs(segments, nodes2)
    seg2run = {}
    for r in runs:
        for s in r["segs"]:
            seg2run[s] = r["run"]
    for s in segments:
        s["run"] = seg2run[s["seg"]]
        for e in entities:
            if e["code"] == s["seg"]:
                e["note"] += f"; 所属{s['run']}"
    d2 = dict(d)
    d2["segments"], d2["nodes"], d2["entities"], d2["runs"] = segments, nodes2, entities, runs
    return d2


def make_report(d, xlsx_path):
    d = _adapt_old(d)
    TOL = d["node_tol"]
    wb = Workbook()
    branches, nodes = d["branches"], d["nodes"]
    segments, entities, runs = d["segments"], d["entities"], d["runs"]
    nmap = {nd["code"]: nd for nd in nodes}

    def xyz_str(p):
        return f'({p[0]:.1f}, {p[1]:.1f}, {p[2]:.1f})'

    # ---- 线段表(核心交付) ----
    rows = [[s["seg"], s["na"], s["nb"], s["length"],
             s["dia"] if s["dia"] is not None else "",
             s["branch"], s["tube"], s["run"],
             s["p0"][0], s["p0"][1], s["p0"][2],
             s["p1"][0], s["p1"][1], s["p1"][2]]
            for s in segments]
    _sheet(wb, "线段表",
           ["线段编码", "起点", "终点", "长度(mm)", "直径(mm)", "对应分支", "来源管体",
            "所属走线", "起点X", "起点Y", "起点Z", "终点X", "终点Y", "终点Z"],
           rows, [10, 10, 10, 11, 10, 16, 26, 10, 11, 11, 11, 11, 11, 11])

    # ---- 节点表 ----
    def nnote(nd):
        if nd["type"] in ("connector", "clamp", "tie"):
            base = "3D实物名见编码对照表"
            sec = nd.get("secondary") or []
            if sec:
                base += "; 同位置设备:" + ",".join(
                    f'{x["code"]}({x["tag"]})' for x in sec)
            return base
        if nd["type"] == "hanging":
            nr = nd.get("nearest")
            if nr:
                return f"悬空; 最近零件{nr['proto']}({nr['tag'].split('/')[-1]}) {nr['dist']}mm"
            return "悬空端"
        if nd["type"] == "fork":
            return "自动分支点(度数>=3)"
        return "连接点"

    rows = []
    for nd in nodes:
        x, y, z = nd["xyz"]
        rows.append([nd["code"], instance_name(nd["name_3d"]), nd["name_3d"],
                     TYPE_CN.get(nd["type"], nd["type"]),
                     round(x, 1), round(y, 1), round(z, 1),
                     nd["degree"], "、".join(nd["segments"]), nnote(nd)])
    _sheet(wb, "节点表",
           ["节点编码", "实例全名", "STEP标签", "类型", "X", "Y", "Z", "度数", "相连线段", "说明"],
           rows, [10, 28, 26, 10, 12, 12, 12, 8, 30, 30])

    # ---- 编码对照表 ----
    def ent_instance(e):
        # 线段: 从note解析对应分支
        if e["category"] == "线段":
            import re
            m = re.search(r"对应分支([^;]+)", e["note"])
            if m:
                return instance_name(m.group(1))
            return e["name_3d"]
        return instance_name(e["name_3d"])
    rows = [[e["code"], e["category"], ent_instance(e), e["name_3d"], e["note"]]
            for e in entities]
    _sheet(wb, "编码对照表", ["编码", "类别", "实例全名", "STEP标签", "说明"],
           rows, [10, 10, 34, 30, 40])

    # ---- 连续走线 ----
    def end_desc(code):
        nd = nmap.get(code)
        if not nd:
            return code
        if nd["type"] == "connector":
            return f'{code}→{nd["name_3d"].split(chr(92))[-1]}'
        if nd["type"] in ("clamp", "tie"):
            return f'{code}→{nd["name_3d"].split(chr(92))[-1]}'
        if nd["type"] == "hanging":
            return f"{code}(悬空)"
        return code

    rows = [[r["run"], " → ".join(r["segs"]), r["length"],
             r["na"], r["nb"], f'{end_desc(r["na"])}；{end_desc(r["nb"])}']
            for r in runs]
    _sheet(wb, "连续走线", ["走线", "线段序列", "总长(mm)", "起点", "终点", "两端说明"],
           rows, [10, 30, 12, 10, 10, 40])

    # ---- 分支清单 ----
    rows = [[b.get("key", b["proto"]), instance_name(b.get("key", "")), b["occ"],
             b["length"], b["dia"],
             len(b.get("segs", [])), "、".join(b.get("segs", [])),
             b.get("node0", ""), b.get("node1", ""),
             xyz_str(b["p0"]), xyz_str(b["p1"]), b["src"]]
            for b in branches]
    _sheet(wb, "分支清单",
           ["分支(STEP)", "实例全名", "实例路径", "中心线长度(mm)", "估算直径(mm)", "线段数",
            "线段编码", "端点A节点", "端点B节点", "端点A坐标", "端点B坐标", "中心线来源"],
           rows, [18, 28, 26, 14, 12, 8, 18, 10, 10, 26, 26, 12])

    # ---- 连接器卡扣清单 ----
    ent_node = {}
    for nd in nodes:
        if not nd.get("entity"):
            continue
        ent_node[nd["entity"]] = nd["code"]
        # 共位合并 tag(如 P#S7+S8) -> 拆回 P#S7, P#S8, 使清单中每个原始实体都能查到编码
        tag = nd["entity"]
        if "#S" in tag and "+" in tag.split("#S")[-1]:
            pre, nums = tag.rsplit("#S", 1)
            for num in nums.split("+S"):
                ent_node[f"{pre}#S{num}"] = nd["code"] + "(合并)"
    rows = [[ent_node.get(c["occ"], ""), instance_name(c["occ"]), c["occ"], c["proto"],
             c["center"][0], c["center"][1], c["center"][2],
             c["bbox"][0], c["bbox"][1], c["bbox"][2]]
            for c in d["connectors"]]
    _sheet(wb, "连接器卡扣清单",
           ["编码", "实例全名", "STEP标签", "产品", "中心X", "中心Y", "中心Z",
            "外形X(mm)", "外形Y(mm)", "外形Z(mm)"],
           rows, [12, 28, 26, 16, 12, 12, 12, 12, 12, 12])

    # ---- 说明 ----
    ws = wb.create_sheet("说明")
    lines = [
        f"{d['file']} 线束 3D 拓扑分析说明(接触式拓扑, 2026-09-27 方案)",
        "",
        "编码规则: 线段 SEG01..(全局); 节点 CON01..(连接器, 只连一个线段端头)/CLP01..(固定卡扣, 路径中间)/BN01..(自动分支点, 度数>=3)/N01..(连接点/悬空端)。",
        "节点命名: 连接器/固定卡扣节点直接使用 3D 实物名; 分支点/普通点自动命名。",
        "编码对照表: 编码 <-> 3D实例全名, 编码为稳定主键; 替换连接器/卡扣时只改对照表, 下游引用编码的行不用动。",
        "线段表的起点/终点为节点编码, 同名即相连, 以此表达拓扑连接关系。",
        "",
        "方法: OpenCASCADE 读取装配树, 管状实体(成组等面积平面端盖+长径比>=4)提取/反推中心线;",
        "实体间最小距离<=2mm 判定接触: 管端-管端=对接, 管端-管侧=T型搭接(向主干投影取站位),",
        "管-实体接触: 单端头接触=连接器(命名节点); 路径穿过/多端头=固定卡扣(命名节点, 劈开主干);",
        f"站位合并: 弧长差与空间距离双条件<=10mm 视为一处搭接; 节点聚类容差 {TOL}mm;",
        "走线在度数!=2 的节点处截断, 不强行穿过分支点配对。",
        "直径: 分支扫掠实体体积 / 中心线长度 → 等效圆直径, 为估算值。",
        "提取端只记录 3D 真值(坐标/弧长/直径), 2D 展开由下游算法处理。",
        "",
        "局限:",
        "- STEP 为哑几何, 不含导线代号/回路/端子型号等电气属性; 电气拓扑需 CATIA EHA 导出 from-to/XML 对齐。",
        "- 平行贴合的多分支(侧-侧接触)未自动拆分, 需手工处理; 接触分类(连接器/扎带)基于尺寸启发式, 请对照 3D 图复核。",
    ]
    for ic in d.get("ignored_contacts", []):
        lines.append(f"- 待人工核对接触: {ic['a']} <-> {ic['b']} (最小距离{ic['dist']}mm): {ic['note']}。")
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t)
        if i == 1:
            ws.cell(row=i, column=1).font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110

    for s in ["Sheet"]:
        if s in wb.sheetnames:
            del wb[s]
    order = ["线段表", "节点表", "编码对照表", "连续走线", "分支清单", "连接器卡扣清单", "说明"]
    for t in order:
        if t in wb.sheetnames:
            wb.move_sheet(t, offset=len(wb.sheetnames))
    wb.active = wb["线段表"]
    wb.save(xlsx_path)


if __name__ == "__main__":
    import json, sys
    d = json.load(open(sys.argv[1] if len(sys.argv) > 1 else "topology.json"))
    make_report(d, "report.xlsx")
    print("saved report.xlsx")
