#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""STEP 线束 3D FROMTO 报表生成器.

输入: CATIA 导出的 STEP (.stp/.step), 哑几何.
输出: 六页 Excel 报表 (页结构对标 3D__fromto_20260122_report 样例):
    Meta / Connectors / Curves / FromToPaths / FromToLatest / Diagnostics

内容规则 (用户确认):
  1. 所有实体一律使用 STEP 文件中的真实名称 (如 Body.190 / 中控大屏 /
     C14\\Tyton-157-00181), 不使用 #Sxx 内部标签.
     名称来源: STEP 文本中 MANIFOLD_SOLID_BREP 出现顺序 == OCP 枚举 solid
     顺序 (X1X 22/22 验证, 见 map_step_names.py); 数量不一致时直接报错退出.
  2. 坐标口径: 一律用线束段端面圆心.
     - 连接器: 所连线束段的端面圆心;
     - 固定卡扣: 先找到接触的线束段, 取最近线束段端面圆心.
  3. 线束段聚类: 端头相接、中间无分支且接头处无固定卡扣 -> 合并为一段;
     接头处有固定卡扣 -> 保留分段点.
  4. 路径: 连接器两两寻径 (Dijkstra, 按 3D 弧长), PathMarker 为
     [连接器, 线段, (卡扣, 线段)*, 连接器] 的 JSON 列表;
     Status = OK / NOT_CONNECTED (断路时在 Diagnostics.BrokenDetails 给出断点).
  5. 若提供 --original 原始 from-to 表, 则只计算该表中的点对并逐行比对
     OrigLength; 否则做连接器全对寻径.

用法:
    python3 fromto_report.py <file.stp> [-o out.xlsx]
        [--original orig.xlsx] [--node-tol 3.0]
"""

import argparse
import csv
import heapq
import json
import math
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# 注意: OCP 相关导入 (harness_topology / map_step_names) 在 main() 内懒加载,
# 纯逻辑函数 (聚类/寻径/报表) 不依赖 OCP, 可独立单元测试.

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SPLIT_TYPES = {"connector", "clamp", "tie"}   # 聚类时的分段点类型
STATUS_OK = "OK"
STATUS_BROKEN = "NOT_CONNECTED"

S_RE = re.compile(r"(?:#|\+)S(\d+)")   # 兼容 P#S7 与 P#S7+S8 两种写法


# ---------------------------------------------------------------- 名称
def build_name_map(step_path):
    """S{i} -> STEP 真实名称. 数量不一致直接报错 (顺序关联的前提)."""
    from map_step_names import parse_brep_names, count_solids_ocp
    names = parse_brep_names(step_path)
    n_solids = count_solids_ocp(step_path)
    if len(names) != n_solids:
        raise RuntimeError(
            "名称映射自检失败: STEP 文本解析到 %d 个顶层 BREP 实体, "
            "OCP 实际枚举到 %d 个 solid, 顺序关联不可信, 终止." % (len(names), n_solids))
    return {"S%d" % i: nm for i, (_sid, nm) in enumerate(names)}


def tag_s_indices(tag):
    return [int(x) for x in S_RE.findall(tag or "")]


def tag_real_names(tag, name_map):
    """tag(如 X#S7+S8) -> 去重后的真实名称列表."""
    out = []
    for si in tag_s_indices(tag):
        nm = name_map.get("S%d" % si, "S%d??" % si)
        if nm not in out:
            out.append(nm)
    return out


def norm_name(s):
    """名称归一化 (对标样例 Diagnostics: 去下划线 + 大写)."""
    return (s or "").replace("_", "").upper()


# ---------------------------------------------------------------- 聚类
def _is_split(node):
    return node["type"] in SPLIT_TYPES or node["degree"] != 2


def cluster_curves(nodes, segments, branch_name_fn):
    """把 SEG 按规则聚成 Curves.

    分段点: 类型为 connector/clamp/tie 的节点, 或度数 != 2 的节点.
    其余度数==2 的普通节点两侧线段合并为一段.
    返回: [{name,length,dia,node_a,node_b,p0,p1,segs:[seg...]}, ...]
    p0/p1 为聚类后两端端面圆心.
    """
    by_code = {n["code"]: n for n in nodes}
    seg_by_id = {s["seg"]: s for s in segments}
    adj = {code: [] for code in by_code}
    for s in segments:
        adj[s["na"]].append(s["seg"])
        adj[s["nb"]].append(s["seg"])

    used = set()
    chains = []  # ([node...], [seg...])

    for code in by_code:
        if not _is_split(by_code[code]):
            continue
        for sid in adj[code]:
            if sid in used:
                continue
            cnodes, csegs = [code], []
            cur_node, cur_seg = code, sid
            while True:
                s = seg_by_id[cur_seg]
                nxt = s["nb"] if s["na"] == cur_node else s["na"]
                csegs.append(cur_seg)
                cnodes.append(nxt)
                used.add(cur_seg)
                if _is_split(by_code[nxt]):
                    break
                rest = [x for x in adj[nxt] if x != cur_seg]
                if len(rest) != 1:
                    break
                cur_node, cur_seg = nxt, rest[0]
            chains.append((cnodes, csegs))

    # 残留: 全由非分段点组成的闭环 (理论上不应出现)
    leftovers = [sid for sid in seg_by_id if sid not in used]
    while leftovers:
        sid0 = leftovers.pop(0)
        s0 = seg_by_id[sid0]
        cnodes, csegs = [s0["na"]], []
        cur_node, cur_seg = s0["na"], sid0
        while True:
            s = seg_by_id[cur_seg]
            nxt = s["nb"] if s["na"] == cur_node else s["na"]
            csegs.append(cur_seg)
            cnodes.append(nxt)
            used.add(cur_seg)
            if cur_seg in leftovers:
                leftovers.remove(cur_seg)
            if nxt == cnodes[0] or _is_split(by_code[nxt]):
                break
            rest = [x for x in adj[nxt] if x != cur_seg and x not in used]
            if not rest:
                break
            cur_node, cur_seg = nxt, rest[0]
        chains.append((cnodes, csegs))

    curves = []
    for cnodes, csegs in chains:
        if not csegs:
            continue
        # 名称: 沿链顺序的管段真实名, 去重相邻重复
        parts = []
        for sid in csegs:
            nm = branch_name_fn(seg_by_id[sid]["branch"])
            if not parts or parts[-1] != nm:
                parts.append(nm)
        name = " + ".join(parts)
        length = round(sum(seg_by_id[s]["length"] for s in csegs), 1)
        tot = sum(seg_by_id[s]["length"] for s in csegs) or 1.0
        dia = round(sum(seg_by_id[s]["length"] * seg_by_id[s].get("dia", 0)
                        for s in csegs) / tot, 2)
        s_first, s_last = seg_by_id[csegs[0]], seg_by_id[csegs[-1]]
        p0 = s_first["p0"] if s_first["na"] == cnodes[0] else s_first["p1"]
        p1 = s_last["p1"] if s_last["nb"] == cnodes[-1] else s_last["p0"]
        curves.append({
            "name": name, "length": length, "dia": dia,
            "node_a": cnodes[0], "node_b": cnodes[-1],
            "p0": [round(float(x), 1) for x in p0],
            "p1": [round(float(x), 1) for x in p1],
            "segs": list(csegs),
        })
    # CurveName 去重
    seen = {}
    for c in curves:
        base = c["name"]
        if base in seen:
            seen[base] += 1
            c["name"] = "%s (%d)" % (base, seen[base])
        else:
            seen[base] = 1
    return curves


# ---------------------------------------------------------------- 坐标
def node_face_center(node, seg_by_id):
    """节点坐标口径: 连接器=所连线段端面圆心; 卡扣=最近线段端面圆心."""
    code = node["code"]
    ends = []
    for sid in node["segments"]:
        s = seg_by_id[sid]
        pt = s["p0"] if s["na"] == code else s["p1"]
        ends.append([float(x) for x in pt])
    if not ends:
        return [round(float(x), 1) for x in node["xyz"]]
    if node["type"] == "connector":
        return [round(x, 1) for x in ends[0]]
    if node["type"] in ("clamp", "tie"):
        xyz = node["xyz"]
        best = min(ends, key=lambda p: sum((p[i] - xyz[i]) ** 2 for i in range(3)))
        return [round(x, 1) for x in best]
    return [round(float(x), 1) for x in node["xyz"]]

# ---------------------------------------------------------------- 连接器
def collect_connectors(nodes, name_map):
    """返回 [{node, partname, name, xyz}]; 含同点次要连接器 (共享节点坐标)."""
    out = []
    for n in nodes:
        if n["type"] != "connector":
            continue
        tags = [n["entity"]] if n.get("entity") else []
        for sec in n.get("secondary", []) or []:
            if sec.get("kind") == "connector":
                tags.append(sec.get("tag"))
        for tag in tags:
            names = tag_real_names(tag, name_map)
            partname = " + ".join(names) if names else (tag or "")
            out.append({"node": n["code"], "partname": partname,
                        "name": partname, "xyz": None, "tag": tag})
    return out


def clamp_partname(node, name_map):
    names = tag_real_names(node.get("entity"), name_map)
    return " + ".join(names) if names else node["code"]


# ---------------------------------------------------------------- 寻径
def build_graph(curves):
    adj = {}
    for ci, c in enumerate(curves):
        adj.setdefault(c["node_a"], []).append((c["node_b"], ci))
        adj.setdefault(c["node_b"], []).append((c["node_a"], ci))
    return adj


def shortest_path(adj, curves, src, dst):
    """Dijkstra. 返回 (node_path, curve_idx_path, length) 或 None."""
    if src == dst:
        return [src], [], 0.0
    dist = {src: 0.0}
    prev = {}
    heap = [(0.0, src)]
    while heap:
        d, u = heapq.heappop(heap)
        if d > dist.get(u, float("inf")):
            continue
        if u == dst:
            break
        for v, ci in adj.get(u, []):
            nd = d + curves[ci]["length"]
            if nd < dist.get(v, float("inf")) - 1e-9:
                dist[v] = nd
                prev[v] = (u, ci)
                heapq.heappush(heap, (nd, v))
    if dst not in prev:
        return None
    node_path, curve_path = [dst], []
    u = dst
    while u != src:
        p, ci = prev[u]
        curve_path.append(ci)
        node_path.append(p)
        u = p
    node_path.reverse()
    curve_path.reverse()
    return node_path, curve_path, round(dist[dst], 1)


def path_marker(conn_a, conn_b, node_path, curve_path, curves, by_code, name_map):
    parts = [conn_a["partname"]]
    for i, ci in enumerate(curve_path):
        parts.append(curves[ci]["name"])
        if i < len(curve_path) - 1:
            mid = by_code[node_path[i + 1]]
            if mid["type"] in ("clamp", "tie"):
                parts.append(clamp_partname(mid, name_map))
    parts.append(conn_b["partname"])
    return json.dumps(parts, ensure_ascii=False)


def component_nodes(adj, start):
    seen = {start}
    stack = [start]
    while stack:
        u = stack.pop()
        for v, _ci in adj.get(u, []):
            if v not in seen:
                seen.add(v)
                stack.append(v)
    return seen


def broken_detail(conn_a, conn_b, adj, curves, by_code):
    """断路诊断: FROM 连通域与 TO 连通域之间最近的两个点 (即断口位置)."""
    comp_a = component_nodes(adj, conn_a["node"])
    comp_b = component_nodes(adj, conn_b["node"])
    best = None
    for u in comp_a:
        pu = by_code[u]["_fx"]
        for v in comp_b:
            pv = by_code[v]["_fx"]
            d = math.sqrt(sum((pu[i] - pv[i]) ** 2 for i in range(3)))
            if best is None or d < best[0]:
                best = (d, u, v)
    if best is None:
        return "无法定位"
    d, u, v = best
    fu = lambda c: "%s [%s]" % (c, ",".join(str(x) for x in by_code[c]["_fx"]))
    return "gap=%.1fmm; 近FROM侧 %s; 近TO侧 %s" % (d, fu(u), fu(v))


# ---------------------------------------------------------------- 原始表
def _match_header(cells):
    low = [(c or "").strip().lower() for c in cells]
    fi = next((i for i, v in enumerate(low) if "from" in v or "起" in v), None)
    ti = next((i for i, v in enumerate(low) if v == "to" or "至" in v or "到" in v), None)
    li = next((i for i, v in enumerate(low)
               if "len" in v or "长" in v or v == "l"), None)
    if fi is None or ti is None:
        return None
    return fi, ti, li


def read_original_table(path):
    """读原始 from-to 表 -> [(from, to, orig_length_or_None)]."""
    rows = []
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig") as f:
            for r in csv.reader(f):
                rows.append(r)
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))
    hdr = None
    for i, r in enumerate(rows[:15]):
        m = _match_header([str(c) if c is not None else "" for c in r])
        if m:
            hdr = (i,) + m
            break
    if hdr is None:
        raise RuntimeError("原始表中未找到 FROM/TO 表头: %s" % path)
    hi, fi, ti, li = hdr
    out = []
    for r in rows[hi + 1:]:
        if fi >= len(r) or ti >= len(r):
            continue
        f = r[fi]
        t = r[ti]
        if f is None or t is None or str(f).strip() == "" or str(t).strip() == "":
            continue
        ol = None
        if li is not None and li < len(r) and r[li] not in (None, ""):
            try:
                ol = float(str(r[li]).replace(",", ""))
            except ValueError:
                ol = None
        out.append((str(f).strip(), str(t).strip(), ol))
    return out
# ---------------------------------------------------------------- 写表
HDR_FONT = Font(bold=True)
HDR_FILL = PatternFill("solid", fgColor="D9E2F3")
CENTER = Alignment(vertical="center")


def _sheet(wb, name, headers, rows, widths, first=False):
    ws = wb.create_sheet(name) if not first else wb.active
    if first:
        ws.title = name
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = HDR_FONT
        c.fill = HDR_FILL
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    return ws


def write_report(out_path, meta_rows, connectors, curves, paths, diag_rows):
    wb = openpyxl.Workbook()
    _sheet(wb, "Meta", ["Key", "Value"], meta_rows, [24, 100], first=True)
    _sheet(wb, "Connectors", ["PartName", "ConnectorName", "X", "Y", "Z"],
           [[c["partname"], c["name"], c["xyz"][0], c["xyz"][1], c["xyz"][2]]
            for c in connectors],
           [42, 42, 12, 12, 12])
    _sheet(wb, "Curves", ["CurveName", "Length(mm)", "ST_X", "ST_Y", "ST_Z",
                           "ED_X", "ED_Y", "ED_Z"],
           [[c["name"], c["length"], c["p0"][0], c["p0"][1], c["p0"][2],
             c["p1"][0], c["p1"][1], c["p1"][2]] for c in curves],
           [48, 12, 10, 10, 10, 10, 10, 10])
    _sheet(wb, "FromToPaths", ["FROM", "TO", "Length(mm)", "PathMarker", "Status"],
           [[p["from"], p["to"], p["length"], p["marker"], p["status"]]
            for p in paths],
           [34, 34, 12, 90, 16])
    _sheet(wb, "FromToLatest",
           ["RowIndex", "FROM", "TO", "OrigLength", "CalcLength(mm)",
            "PathMarker", "Status"],
           [[p["row_idx"], p["from"], p["to"], p["orig_len"], p["length"],
             p["marker"], p["status"]] for p in paths],
           [10, 34, 34, 12, 14, 90, 16])
    _sheet(wb, "Diagnostics", ["Section", "Key", "Value"], diag_rows,
           [16, 56, 80])
    wb.save(out_path)


# ---------------------------------------------------------------- 主流程
def main():
    ap = argparse.ArgumentParser(description="STEP 线束 3D FROMTO 报表")
    ap.add_argument("step", help="输入 STEP 文件 (.stp/.step)")
    ap.add_argument("-o", "--out", default=None, help="输出 xlsx 路径")
    ap.add_argument("--original", default=None, help="原始 from-to 表 (xlsx/csv, 可选)")
    ap.add_argument("--node-tol", type=float, default=3.0, help="节点容差 mm")
    args = ap.parse_args()

    from harness_topology import analyze

    step_path = args.step
    ts = time.strftime("%Y%m%d_%H%M%S")
    stem = os.path.splitext(os.path.basename(step_path))[0]
    out_path = args.out or "%s_fromto_report_%s.xlsx" % (stem, ts)

    # 名称映射自检前置: 失败时直接退出, 不浪费拓扑解析时间
    print("[0/6] 真实名称映射自检 ...", flush=True)
    name_map = build_name_map(step_path)
    print("      %d 个实体名称, 自检通过" % len(name_map), flush=True)

    print("[1/6] 拓扑解析 ...", flush=True)
    topo = analyze(step_path, tol=args.node_tol)
    nodes, segments = topo["nodes"], topo["segments"]
    by_code = {n["code"]: n for n in nodes}
    seg_by_id = {s["seg"]: s for s in segments}

    print("[2/6] 名称关联 ...", flush=True)

    def branch_name(bkey):
        names = tag_real_names(bkey, name_map)
        return " + ".join(names) if names else bkey

    print("[3/6] 线段聚类 ...", flush=True)
    curves = cluster_curves(nodes, segments, branch_name)
    print("      %d SEG -> %d Curves" % (len(segments), len(curves)), flush=True)

    print("[4/6] 坐标口径 (端面圆心) ...", flush=True)
    for n in nodes:
        n["_fx"] = node_face_center(n, seg_by_id)
        if n["type"] in ("clamp", "tie"):
            n["_partname"] = clamp_partname(n, name_map)
    connectors = collect_connectors(nodes, name_map)
    for c in connectors:
        c["xyz"] = by_code[c["node"]]["_fx"]
    print("      %d 连接器" % len(connectors), flush=True)

    print("[5/6] 寻径 ...", flush=True)
    adj = build_graph(curves)
    orig_rows = read_original_table(args.original) if args.original else None
    by_norm = {}
    for c in connectors:
        by_norm.setdefault(norm_name(c["name"]), []).append(c)

    pairs = []       # (conn_a, conn_b, orig_len)
    diag = []
    if orig_rows:
        for i, (of, ot, ol) in enumerate(orig_rows, start=1):
            ca = by_norm.get(norm_name(of), [None])[0]
            cb = by_norm.get(norm_name(ot), [None])[0]
            if ca and cb:
                pairs.append((ca, cb, ol))
                diag.append(("FromToMatch", "orig row %d: %s -> %s" % (i, of, ot),
                             "MATCHED"))
            else:
                miss = []
                if not ca:
                    miss.append("FROM '%s'" % of)
                if not cb:
                    miss.append("TO '%s'" % ot)
                diag.append(("FromToMatch", "orig row %d: %s -> %s" % (i, of, ot),
                             "UNMATCHED (%s)" % ", ".join(miss)))
    else:
        for i in range(len(connectors)):
            for j in range(i + 1, len(connectors)):
                pairs.append((connectors[i], connectors[j], None))

    paths = []
    n_broken = 0
    for ca, cb, ol in pairs:
        r = shortest_path(adj, curves, ca["node"], cb["node"])
        if r is None:
            paths.append({"from": ca["name"], "to": cb["name"], "length": None,
                          "marker": None, "status": STATUS_BROKEN,
                          "orig_len": ol, "row_idx": len(paths) + 2,
                          "ca": ca, "cb": cb})
            n_broken += 1
        else:
            node_path, curve_path, length = r
            marker = path_marker(ca, cb, node_path, curve_path, curves,
                                 by_code, name_map)
            paths.append({"from": ca["name"], "to": cb["name"],
                          "length": length, "marker": marker,
                          "status": STATUS_OK, "orig_len": ol,
                          "row_idx": len(paths) + 2, "ca": ca, "cb": cb})
    print("      %d 对: OK=%d NOT_CONNECTED=%d"
          % (len(paths), len(paths) - n_broken, n_broken), flush=True)

    print("[6/6] 写报表 ...", flush=True)
    diag_counts = [
        ("Counts", "connectors", len(connectors)),
        ("Counts", "curves", len(curves)),
        ("Counts", "nodes", len(nodes)),
        ("Counts", "edges", len(curves)),
        ("Counts", "fromto_rows", len(paths)),
        ("Counts", "not_connected", n_broken),
    ]
    diag_norm = [("Connectors", c["partname"], norm_name(c["name"]))
                 for c in connectors]
    seen_norm = {}
    for c in connectors:
        seen_norm.setdefault(norm_name(c["name"]), []).append(c["partname"])
    diag_coll = [("NameCollisions", k, "%d: %s" % (len(v), " | ".join(v)))
                 for k, v in sorted(seen_norm.items()) if len(v) > 1]
    diag_broken = []
    for p in paths:
        if p["status"] == STATUS_BROKEN:
            diag_broken.append(
                ("BrokenDetails", "%s -> %s" % (p["from"], p["to"]),
                 broken_detail(p["ca"], p["cb"], adj, curves, by_code)))
    diag_rows = diag_counts + [("Counts", "----", "----")] + diag_norm + \
        diag_coll + diag + diag_broken

    meta_rows = [
        ("ReportCreatedAt", ts),
        ("Tool", "harness_3d/fromto_report.py"),
        ("InputStep", os.path.abspath(step_path)),
        ("NodeTolMm", args.node_tol),
        ("NameMap", "S{i} = SHAPE_REPRESENTATION 第 i 个 item 的 BREP 实体名称"
                    "(含 BREP_WITH_VOIDS); %d solids, 自检通过" % len(name_map)),
        ("OriginalTable", os.path.abspath(args.original) if args.original
         else "(none: connector all-pairs)"),
        ("CoordinateRule", "连接器/卡扣坐标 = 所连线束段端面圆心"),
        ("ClusterRule", "端头相接且无分支无卡扣则合并; 卡扣处保留分段点"),
        ("StatusValues", "OK=连通 / NOT_CONNECTED=断路(见Diagnostics.BrokenDetails)"),
        ("OutputPath", os.path.abspath(out_path)),
    ]
    write_report(out_path, meta_rows, connectors, curves, paths, diag_rows)
    print("已生成: %s" % os.path.abspath(out_path), flush=True)


if __name__ == "__main__":
    main()
