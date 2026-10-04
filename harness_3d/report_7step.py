#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""7步法格式报表: 从 topology.json + 名称映射生成 CODEX 7步表的7个sheet.
用法: python3 report_7step.py <topology.json> <name_map.json> -o <out.xlsx>
"""
import argparse
import json
import math
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict


def load_name_map(path):
    mp = {}
    for e in json.load(open(path, encoding="utf-8")):
        mp[e["stag"]] = e["name"]
    return mp


def real_names(tag, mp):
    """IP_AllCATPart#S53+55+57+105 -> 真名列表."""
    m = re.search(r"#(.+)$", tag)
    if not m:
        return [tag]
    out = []
    for p in m.group(1).split("+"):
        mm = re.match(r"S(\d+)", p)
        out.append(mp.get("S" + mm.group(1), p) if mm else p)
    return out


def real_name_joined(tag, mp, sep="+"):
    return sep.join(real_names(tag, mp))


def r1(x):
    return round(float(x), 1)


def straight_curved_split(polyline):
    """估算圆柱(直)长度 vs 弯曲长度: 按转向角划分."""
    if len(polyline) < 3:
        L = sum(math.dist(polyline[i], polyline[i+1]) for i in range(len(polyline)-1))
        return r1(L), 0.0
    straight, curved = 0.0, 0.0
    for i in range(len(polyline) - 2):
        a = polyline[i]; b = polyline[i+1]; c = polyline[i+2]
        v1 = [b[k]-a[k] for k in range(3)]
        v2 = [c[k]-b[k] for k in range(3)]
        l1 = math.dist(a, b); l2 = math.dist(b, c)
        if l1 < 1e-9 or l2 < 1e-9:
            straight += l1
            continue
        cosang = sum(v1[k]*v2[k] for k in range(3)) / (l1*l2)
        cosang = max(-1.0, min(1.0, cosang))
        ang = math.degrees(math.acos(cosang))
        if ang < 8.0:
            straight += l1
        else:
            curved += l1
    # 最后一段
    l_last = math.dist(polyline[-2], polyline[-1])
    straight += l_last
    return r1(straight), r1(curved)


def style_header(ws, ncols):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF", size=11)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    widths = [10, 34, 14, 14, 14, 12, 12, 12, 12, 12, 12, 12, 14]
    for i, w in enumerate(widths[:ncols], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("topo", help="topology.json")
    ap.add_argument("namemap", help="step_name_map.json")
    ap.add_argument("-o", "--out", required=True)
    args = ap.parse_args()

    topo = json.load(open(args.topo, encoding="utf-8"))
    mp = load_name_map(args.namemap)
    branches = topo["branches"]
    nodes = topo["nodes"]
    segs = topo["segments"]
    relations = topo.get("relations", {})
    devices = topo.get("connectors", [])  # 全部设备(含未接触)
    nc = {n["code"]: n for n in nodes}
    bkey2idx = {b["key"]: i for i, b in enumerate(branches)}

    wb = openpyxl.Workbook()

    # ---- Sheet1: 管状实体 ----
    ws = wb.active
    ws.title = "1-管状实体"
    ws.append(["序号", "实体名称", "总长度(mm)", "圆柱长度(mm)", "弯曲长度(mm)",
               "半径(mm)", "起点X", "起点Y", "起点Z", "终点X", "终点Y", "终点Z"])
    for i, b in enumerate(branches, 1):
        pl = b.get("polyline") or [b["p0"], b["p1"]]
        sl, cl = straight_curved_split(pl)
        ws.append([i, real_name_joined(b["key"], mp), b["length"],
                   sl, cl, round(b["dia"]/2, 2) if b["dia"] else "",
                   r1(b["p0"][0]), r1(b["p0"][1]), r1(b["p0"][2]),
                   r1(b["p1"][0]), r1(b["p1"][1]), r1(b["p1"][2])])
    style_header(ws, 12)

    # ---- Sheet2: 接触关系 ----
    ws = wb.create_sheet("2-接触关系")
    ws.append(["序号", "实体名称", "实体中心X", "实体中心Y", "实体中心Z",
               "接触管状实体", "接触端点", "接触点X", "接触点Y", "接触点Z", "距离(mm)"])
    terms_by_tag = defaultdict(list)
    for t in relations.get("terminals", []):
        terms_by_tag[t["tag"]].append(t)
    ties_by_tag = defaultdict(list)
    for t in relations.get("tie_stations", []):
        ties_by_tag[t["tag"]].append(t)
    row = 0
    for dev in devices:
        tag = dev["occ"]
        rname = real_name_joined(tag, mp)
        cx, cy, cz = dev["center"]
        contacts = []
        for t in terms_by_tag.get(tag, []):
            b = branches[t["branch"]]
            exyz = b["p0"] if t["end"] == 0 else b["p1"]
            contacts.append((real_name_joined(b["key"], mp),
                             f"p{t['end']}", exyz, t.get("dist", 0.0)))
        for t in ties_by_tag.get(tag, []):
            b = branches[t["branch"]]
            contacts.append((real_name_joined(b["key"], mp),
                             f"中部{t['s']:.0f}mm", t["xyz"], t.get("dev", 0.0)))
        if not contacts:
            row += 1
            ws.append([row, rname, r1(cx), r1(cy), r1(cz),
                       "无接触", "", "", "", "", ""])
        for bname, pend, pxyz, dist in contacts:
            row += 1
            ws.append([row, rname, r1(cx), r1(cy), r1(cz),
                       bname, pend, r1(pxyz[0]), r1(pxyz[1]), r1(pxyz[2]),
                       round(dist, 1)])
    style_header(ws, 11)

    # ---- Sheet3: 定位坐标 ----
    ws = wb.create_sheet("3-定位坐标")
    ws.append(["序号", "实体名称", "定位坐标X", "定位坐标Y", "定位坐标Z",
               "接触数", "推理原因"])
    # 设备 -> 节点
    node_by_tag = {}
    for n in nodes:
        if n.get("name_3d"):
            node_by_tag[n["name_3d"]] = n
    row = 0
    for dev in devices:
        tag = dev["occ"]
        rname = real_name_joined(tag, mp)
        n = node_by_tag.get(tag)
        nt = len(terms_by_tag.get(tag, [])) + len(ties_by_tag.get(tag, []))
        if n:
            xyz = n["xyz"]
            reasons = []
            for t in terms_by_tag.get(tag, []):
                b = branches[t["branch"]]
                reasons.append(f"端部接触:{real_name_joined(b['key'],mp)}.p{t['end']}")
            for t in ties_by_tag.get(tag, []):
                b = branches[t["branch"]]
                reasons.append(f"中部固定:{real_name_joined(b['key'],mp)}@{t['s']:.0f}mm")
            reason = "; ".join(reasons)
        else:
            xyz = dev["center"]
            reason = "无接触，使用原始中心"
        row += 1
        ws.append([row, rname, r1(xyz[0]), r1(xyz[1]), r1(xyz[2]), nt, reason])
    style_header(ws, 7)

    # ---- Sheet4: 分段点 ----
    ws = wb.create_sheet("4-分段点")
    ws.append(["分段点ID", "类型", "关联实体", "位置X", "位置Y", "位置Z", "备注"])
    sp_id = 0
    for t in relations.get("tie_stations", []):
        sp_id += 1
        b = branches[t["branch"]]
        ws.append([f"SP{sp_id:03d}", "卡扣", real_name_joined(t["tag"], mp),
                   r1(t["xyz"][0]), r1(t["xyz"][1]), r1(t["xyz"][2]),
                   f"中部分段: {real_name_joined(b['key'],mp)}@{t['s']:.1f}mm"])
    for t in relations.get("taps", []):
        sp_id += 1
        b = branches[t["main"]]
        tb = branches[t["tap"]]
        ws.append([f"SP{sp_id:03d}", "搭接", real_name_joined(tb["key"], mp),
                   r1(t["xyz"][0]), r1(t["xyz"][1]), r1(t["xyz"][2]),
                   f"搭接分段: {real_name_joined(tb['key'],mp)}->{real_name_joined(b['key'],mp)}@{t['s']:.1f}mm"])
    style_header(ws, 7)

    # ---- Sheet5: 合并决策 ----
    ws = wb.create_sheet("5-合并决策")
    ws.append(["序号", "管状实体A", "端点A", "管状实体B", "端点B",
               "端点距离(mm)", "是否有分段点", "分段实体", "决策", "原因"])
    # 端头有设备的集合
    term_ends = set((t["branch"], t["end"]) for t in relations.get("terminals", []))
    tag_at = {}
    for t in relations.get("terminals", []):
        tag_at[(t["branch"], t["end"])] = t["tag"]
    row = 0
    for (i, ei, j, ej) in relations.get("end_ends", []):
        bi, bj = branches[i], branches[j]
        pei = bi["p0"] if ei == 0 else bi["p1"]
        pej = bj["p0"] if ej == 0 else bj["p1"]
        dist = math.dist(pei, pej)
        has_dev = (i, ei) in term_ends or (j, ej) in term_ends
        dev_tag = tag_at.get((i, ei), tag_at.get((j, ej), ""))
        dev_name = real_name_joined(dev_tag, mp) if dev_tag else ""
        if has_dev:
            decision, reason = "NOT_MERGE", f"接头处有设备: {dev_name}"
        else:
            decision, reason = "MERGE", "端-端直接对接，无设备"
        row += 1
        ws.append([row, real_name_joined(bi["key"], mp), f"p{ei}",
                   real_name_joined(bj["key"], mp), f"p{ej}",
                   round(dist, 1), "是" if has_dev else "否", dev_name,
                   decision, reason])
    style_header(ws, 10)

    # ---- Sheet6: 拓扑关系 ----
    ws = wb.create_sheet("6-拓扑关系")
    ws.append(["节点ID", "位置X", "位置Y", "位置Z", "成员数", "连接线束段"])
    for n in nodes:
        ws.append([n["code"], r1(n["xyz"][0]), r1(n["xyz"][1]), r1(n["xyz"][2]),
                   n["degree"], "; ".join(n.get("segments", []))])
    style_header(ws, 6)

    # ---- Sheet7: 实体定位 ----
    ws = wb.create_sheet("7-实体定位")
    ws.append(["序号", "实体名称", "类型", "定位坐标X", "定位坐标Y", "定位坐标Z",
               "挂在线束段", "参数t", "距起点距离(mm)", "距线束段距离(mm)"])
    tname = {"connector": "connector", "clamp": "clamp", "tie": "tie"}
    row = 0
    # 从 tie_stations / terminals 反推每个设备的定位
    seen_tag = set()
    for t in relations.get("tie_stations", []):
        tag = t["tag"]
        if tag in seen_tag:
            continue
        seen_tag.add(tag)
        b = branches[t["branch"]]
        L = b["length"] or 1
        s = t["s"]
        row += 1
        ws.append([row, real_name_joined(tag, mp), t.get("kind", "clamp"),
                   r1(t["xyz"][0]), r1(t["xyz"][1]), r1(t["xyz"][2]),
                   real_name_joined(b["key"], mp), round(s/L, 3),
                   round(s, 1), round(t.get("dev", 0.0), 1)])
    for t in relations.get("terminals", []):
        tag = t["tag"]
        if tag in seen_tag:
            continue
        seen_tag.add(tag)
        b = branches[t["branch"]]
        L = b["length"] or 1
        s = 0.0 if t["end"] == 0 else L
        exyz = b["p0"] if t["end"] == 0 else b["p1"]
        row += 1
        ws.append([row, real_name_joined(tag, mp), t.get("kind", "connector"),
                   r1(exyz[0]), r1(exyz[1]), r1(exyz[2]),
                   real_name_joined(b["key"], mp), round(s/L, 3),
                   round(s, 1), 0.0])
    # 无接触设备
    for dev in devices:
        if dev["occ"] in seen_tag:
            continue
        seen_tag.add(dev["occ"])
        cx, cy, cz = dev["center"]
        row += 1
        ws.append([row, real_name_joined(dev["occ"], mp), "uncontacted",
                   r1(cx), r1(cy), r1(cz), "", "", "", ""])
    style_header(ws, 10)

    wb.save(args.out)
    print(f"已生成: {args.out}")


if __name__ == "__main__":
    main()
